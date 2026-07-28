from __future__ import annotations

import hashlib
import logging
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from .db import Database
from .mailer import send_html_email
from .mail_templates import ensure_mail_templates, render_mail_template
from .settings import Settings
from .test_catalog import load_test_definitions

LOGGER = logging.getLogger("invite-mailer.technical-alerts")


def utc_now() -> datetime:
    return datetime.now(timezone.utc).replace(microsecond=0)


def _read_setting(connection, key: str, default: str = "") -> str:
    row = connection.execute("SELECT value FROM app_settings WHERE key = ?", (key,)).fetchone()
    return str(row["value"]) if row else default


def technical_settings(db: Database) -> dict:
    with db.connect() as connection:
        try:
            repeat_hours = max(1, int(_read_setting(connection, "technical_repeat_hours", "72")))
        except ValueError:
            repeat_hours = 72
        recipients = [
            {"name": str(row["name"]), "email": str(row["email"]).strip()}
            for row in connection.execute("""
                SELECT name, email FROM reviewers
                WHERE enabled = 1 AND receives_technical_errors = 1
                ORDER BY name COLLATE NOCASE, email COLLATE NOCASE
            """).fetchall() if str(row["email"]).strip()
        ]
    return {"recipients": recipients, "repeat_hours": repeat_hours}


def _fingerprint(error_type: str, employee_email: str, error_text: str) -> str:
    value = "\n".join((error_type.strip(), employee_email.strip().lower(), error_text.strip()))
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def notify_technical_error(settings: Settings, db: Database, *, subject: str, fio: str,
                           email_address: str, error_type: str, error_text: str = "") -> bool:
    """Регистрирует ошибку. Общая сводка отправляется после завершения рассылки."""
    config = technical_settings(db)
    now = utc_now()
    fingerprint = _fingerprint(error_type, email_address, error_text)
    with db.connect() as connection:
        previous = connection.execute("""
            SELECT detected_at, notified_at FROM technical_errors
            WHERE fingerprint = ? ORDER BY id DESC LIMIT 1
        """, (fingerprint,)).fetchone()
        if previous:
            stamp = previous["notified_at"] or previous["detected_at"]
            try:
                last = datetime.fromisoformat(stamp)
                if last.tzinfo is None:
                    last = last.replace(tzinfo=timezone.utc)
                if now - last < timedelta(hours=config["repeat_hours"]):
                    return False
            except (ValueError, TypeError):
                pass
        connection.execute("""
            INSERT INTO technical_errors(
                fingerprint, fio, email, error_type, error_text,
                detected_at, notified_at, notification_error
            ) VALUES (?, ?, ?, ?, ?, ?, NULL, NULL)
        """, (fingerprint, fio, email_address, error_type, error_text, now.isoformat()))
    return True


def _build_xlsx(rows) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Ошибки отправки"
    headers = ["Дата", "ФИО", "E-mail", "Тип ошибки", "Текст ошибки"]
    ws.append(headers)
    for cell in ws[1]: cell.font = Font(bold=True)
    for row in rows:
        try:
            dt = datetime.fromisoformat(row["detected_at"]).astimezone().strftime("%d.%m.%Y %H:%M:%S")
        except Exception:
            dt = str(row["detected_at"])
        ws.append([dt, row["fio"] or "", row["email"] or "", row["error_type"] or "", row["error_text"] or ""])
    widths = [20, 38, 34, 48, 80]
    for i, width in enumerate(widths, 1): ws.column_dimensions[chr(64+i)].width = width
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    out = BytesIO(); wb.save(out); return out.getvalue()


def dispatch_technical_error_digest(settings: Settings, db: Database) -> dict:
    """Отправляет одно письмо с XLSX-файлом по всем новым ошибкам текущей рассылки."""
    config = technical_settings(db)
    if not config["recipients"]:
        return {"errors": 0, "delivered": 0}
    with db.connect() as connection:
        rows = connection.execute("""
            SELECT * FROM technical_errors
            WHERE notified_at IS NULL AND (notification_error IS NULL OR notification_error = '')
            ORDER BY detected_at, id
        """).fetchall()
    if not rows:
        return {"errors": 0, "delivered": 0}

    settings.templates[:] = load_test_definitions(settings, db)
    ensure_mail_templates(db, settings.templates)
    context = {
        "subject": "Ошибки отправки писем",
        "fio": "Сводный отчет",
        "email": "См. вложение",
        "error_type": f"Обнаружено ошибок: {len(rows)}",
        "error_text": "Подробный список ошибок находится во вложенном файле.",
        "detected_at": utc_now().astimezone().strftime("%d.%m.%Y %H:%M"),
    }
    subject, body, enabled = render_mail_template(db, "technical", "*", context)
    if not enabled:
        return {"errors": len(rows), "delivered": 0}
    if not subject.strip(): subject = "Ошибки отправки писем"
    if not body.strip(): body = f"<p>Во время рассылки обнаружено ошибок: {len(rows)}.</p><p>Подробности находятся во вложенном файле.</p>"
    filename = f"mail-errors-{utc_now().astimezone().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
    attachment = (filename, _build_xlsx(rows), "application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    delivered = 0; failures = []
    for recipient in config["recipients"]:
        try:
            send_html_email(settings.smtp, recipient["email"], subject, body, attachments=[attachment])
            delivered += 1
        except Exception as error:
            LOGGER.exception("Не удалось отправить сводку технических ошибок на %s", recipient["email"])
            failures.append(f"{recipient['name']} <{recipient['email']}>: {error}")
    now_iso = utc_now().isoformat(); ids = [int(r["id"]) for r in rows]
    with db.connect() as connection:
        placeholders = ",".join("?" for _ in ids)
        if delivered:
            connection.execute(f"UPDATE technical_errors SET notified_at=?, notification_error=NULL WHERE id IN ({placeholders})", (now_iso, *ids))
        elif failures:
            connection.execute(f"UPDATE technical_errors SET notification_error=? WHERE id IN ({placeholders})", ("\n".join(failures), *ids))
    return {"errors": len(rows), "delivered": delivered}
